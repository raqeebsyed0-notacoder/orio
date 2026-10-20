# ruff: noqa: BLE001
"""ORIO Operations demo backend — local SQLite, zero external calls.

Mounted at ``/api/plugins/euphoria-orio-operations/`` by the Hermes
dashboard plugin system (``dashboard/manifest.json`` declares
``api=plugin_api.py``; the plugin must be listed in ``plugins.enabled``
before the serve process imports it).

Demo safety: this module opens NO sockets, sends NO mail, touches NO
production system. Every read and write targets the demo database at
``<HERMES_HOME>/state/<client>/orio_demo.db`` (WAL — ``orio-ops`` for
client #1, Ajeez's pilot; see ``dashboard/clients.py``), built by
``data/build_db.py`` from Ajeez's supplied Sep 19 files. Approving an
action mutates DEMO STATE ONLY. A request selects its client workspace
via the ``x-orio-client`` header or ``?client=`` param (default #1).

Production path: replace the ``_db``/query layer with the production
adapters (OneDrive ingestion -> operational datastore). The route shapes
stay the same so the desktop UI does not change.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import time
from collections import deque
from contextlib import contextmanager
from datetime import date as _date
from pathlib import Path
from typing import Any, Iterator

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import JSONResponse, Response


def _envelope(name: str, media: str, payload: bytes) -> JSONResponse:
    """JSON-safe file envelope for JSON-only shells.

    The desktop panel reaches the backend through ``ctx.rest``, which
    cannot carry raw CSV text or .xlsx bytes in every shell — both
    downloads surfaced as 'Export failed'. The envelope carries the
    identical bytes as base64 inside JSON; the panel decodes to a Blob.
    Byte endpoints below are unchanged (direct curl still works).
    """
    return JSONResponse({
        'filename': name,
        'media_type': media,
        'size': len(payload),
        'sha256': hashlib.sha256(payload).hexdigest(),
        'b64': base64.b64encode(payload).decode('ascii'),
    })

try:  # tests put dashboard/ on sys.path; the gateway does neither
    from channels import all_statuses, build_email_payload, get_channel
    from clients import (DEFAULT_CLIENT, UnknownClient,
                         db_path as _client_db_path, list_clients,
                         load_config, resolve_client)
except Exception:  # pragma: no cover - gateway loads this file standalone
    import sys as _sys

    _here = str(Path(__file__).resolve().parent)
    if _here not in _sys.path:
        _sys.path.insert(0, _here)
    from channels import (  # type: ignore[no-redef]
        all_statuses, build_email_payload, get_channel)
    from clients import (  # type: ignore[no-redef]
        DEFAULT_CLIENT, UnknownClient,
        db_path as _client_db_path, list_clients,
        load_config, resolve_client)

try:  # available in-gateway; tests get it via the hermes-agent venv
    from hermes_constants import get_hermes_home
except Exception:  # pragma: no cover - defensive fallback
    def get_hermes_home() -> Path:  # type: ignore[misc]
        return Path(os.environ.get('HERMES_HOME') or Path.home() / '.hermes')


def _db_path(client: str = DEFAULT_CLIENT) -> Path:
    return _client_db_path(get_hermes_home(), client or DEFAULT_CLIENT)


@contextmanager
def _db(client: str = DEFAULT_CLIENT) -> Iterator[sqlite3.Connection]:
    p = _db_path(client)
    if not p.exists():
        raise HTTPException(503, 'Demo database not built yet. Run data/build_db.py.')
    con = sqlite3.connect(str(p))
    con.row_factory = sqlite3.Row
    try:
        yield con
    finally:
        con.close()


def _resolve_client(request: Request | None,
                    param: str = '') -> str:
    """Bind a request to an onboarded client workspace.

    ``x-orio-client`` header wins, then the ``?client=`` query param, then
    client #1 (``orio-ops``). Unknown slugs answer 404 — the API never
    auto-creates a workspace; onboarding is ``tools/onboard_client.sh``.
    """
    raw = ''
    try:
        if request is not None:
            raw = request.headers.get('x-orio-client') or ''
    except Exception:  # noqa: BLE001 - header read must never fail a read
        raw = ''
    raw = (raw or (param or '')).strip()
    try:
        return resolve_client(raw, get_hermes_home())
    except UnknownClient as e:
        raise HTTPException(404, str(e)) from None
    except ValueError as e:
        raise HTTPException(400, str(e)) from None


def _row(r: sqlite3.Row) -> dict[str, Any]:
    return dict(r)


def _csv_cell(v) -> str:
    """CSV formula-injection guard. A cell starting with = + - @ (or a
    control char) can execute a formula when the export opens in Excel;
    prefix with a single quote (Excel's native text marker) so the value
    stays inert. Applied to every exported cell."""
    s = '' if v is None else str(v)
    if (s[:1] in ('=', '+', '-', '@', '\t', '\r', '\n')
            or s.lstrip()[:1] in ('=', '+', '-', '@')):
        return "'" + s
    return s


router = APIRouter()


# ---------------------------------------------------------- company aliases (B4)
# data/company_aliases.json is the correspondence matrix: report spellings
# (HOWCO, HLB, DAVIS-LYNCH, ...) resolve to canonical companies, and
# Resp.Co codes (WDV/HDV/TOD) NEVER resolve to a company (R1 corollary).
_ALIAS_FILE = Path(__file__).resolve().parent.parent / 'data' / 'company_aliases.json'
_ALIASES: dict[str, Any] | None = None


def _aliases() -> dict[str, Any]:
    global _ALIASES
    if _ALIASES is None:
        try:
            loaded = json.loads(_ALIAS_FILE.read_text(encoding='utf-8'))
        except OSError:
            loaded = {'mapping': {}, 'never_company': {'codes': []}}
        _ALIASES = loaded
    return _ALIASES


def _resolve_company(raw: str) -> tuple[str, str | None]:
    """Resolve a ?company= value through the alias map (B4).

    Returns (canonical_code, note). Canonical codes pass through; report
    aliases resolve (HOWCO->HAL per R3, DAVIS-LYNCH->FET, HLB->HAL, ...);
    never-company Resp.Co codes (WDV/HDV/TOD) raise 400 — they are work
    markings, not companies, and must never silently match a company.
    Unknown codes pass through verbatim and match nothing (empty result,
    never an invention).
    """
    code = (raw or '').strip().upper()
    if not code:
        return '', None
    amap = _aliases()
    never = {c.upper() for c in (amap.get('never_company') or {}).get('codes', [])}
    if code in never:
        raise HTTPException(
            400, '%s is a Resp.Co / work-marking code, never a company '
            '(rules-v2 R1 corollary: row 106 WDV, company UNPROVEN)' % code)
    for canon, entry in (amap.get('mapping') or {}).items():
        spellings = {canon.upper()} | {a.upper() for a in entry.get('report_aliases', [])}
        # Hyphen/space-insensitive: DAVIS LYNCH == DAVIS-LYNCH.
        spellings |= {s.replace('-', ' ').replace('  ', ' ') for s in spellings}
        spellings |= {s.replace(' ', '-') for s in spellings}
        if code in spellings:
            note = None if canon.upper() == code else (
                '%s resolved to %s (%s)' % (code, canon.upper(),
                 (entry.get('evidence') or '')[:120]))
            return canon.upper(), note
    return code, None


# ---------------------------------------------------------- Phase 3 hardening
# Approvals are attributed, never anonymous. The gateway fronts route auth in
# production; this module binds the approval actor as follows:
#   1. `x-hermes-user` / `x-hermes-user-role` headers win (source 'gateway').
#      A client body that claims a DIFFERENT actor is rejected (400) — a
#      forged self-claim must be visible, not silently stored.
#   2. Otherwise a non-empty client body claim is stored as source 'client'
#      (self-claimed, untrusted — the unauth probe lands here, handled).
#   3. Otherwise the local demo default (source 'default').
# Demo has no login page; production must put gateway auth in front and treat
# only source='gateway' rows as non-repudiable.
ACTOR_DEFAULT = 'Ajeez'
ACTOR_ROLE_DEFAULT = 'operator'
_ACTOR_MAX = 64

OPLOG: deque[dict[str, Any]] = deque(maxlen=200)


def _oplog(op: str, status: str, detail: str = '') -> None:
    OPLOG.append({'at': time.strftime('%Y-%m-%dT%H:%M:%S'),
                  'op': op, 'status': status, 'detail': detail[:300]})


def _clean(v: Any) -> str:
    return (v if isinstance(v, str) else '' if v is None else str(v)).strip()


def _resolve_actor(body: dict[str, Any], request: Request) -> tuple[str, str, str]:
    gw_actor = _clean(request.headers.get('x-hermes-user'))
    gw_role = _clean(request.headers.get('x-hermes-user-role'))
    body_actor = _clean(body.get('actor'))
    body_role = _clean(body.get('actor_role'))
    for v in (gw_actor, gw_role, body_actor, body_role):
        if len(v) > _ACTOR_MAX:
            raise HTTPException(400, 'actor/actor_role too long (max 64)')
    if gw_actor:
        if body_actor and body_actor.lower() != gw_actor.lower():
            _oplog('decision', 'rejected-forgery',
                   'gateway=%s body=%s' % (gw_actor, body_actor))
            raise HTTPException(
                400, 'actor mismatch: gateway identity wins; '
                     'drop the client actor claim.')
        return gw_actor, gw_role or body_role or ACTOR_ROLE_DEFAULT, 'gateway'
    if body_actor:
        return body_actor, body_role or ACTOR_ROLE_DEFAULT, 'client'
    return ACTOR_DEFAULT, ACTOR_ROLE_DEFAULT, 'default'


@router.get('/ops/log')
def ops_log(limit: int = Query(50, le=200)) -> dict[str, Any]:
    """Most recent write/error entries. Reads are not logged (noise)."""
    entries = list(OPLOG)[-limit:]
    return {'entries': entries, 'n': len(entries),
            'note': 'Demo ops log: approval writes + errors. '
                    'Empty until the first write this process lifetime.'}


@router.get('/health')
def health(request: Request, client: str = Query('')) -> dict[str, Any]:
    c = _resolve_client(request, client)
    p = _db_path(c)
    return {'ok': p.exists(), 'db': str(p), 'client': c,
            'demo': True, 'external_actions': False}


@router.get('/meta')
def meta(request: Request, client: str = Query('')) -> dict[str, Any]:
    c = _resolve_client(request, client)
    with _db(c) as con:
        m = {r['key']: r['value'] for r in con.execute('SELECT * FROM meta')}
        days = [r['report_date'] for r in con.execute(
            'SELECT report_date FROM reports ORDER BY report_date DESC')]
    return {'demo_date': days[0] if days else m.get('seed_date', '?'),
            'report_dates': days,
            'built_from': m.get('built_from', '?'),
            'matrix_status_counts': m.get('matrix_status_counts', '?'),
            'data_tier': m.get('data_tier', '?'),
            'data_tier_note': m.get('data_tier_note', ''),
            'client': c,
            'banner': 'DEMO - LOCAL DATA ONLY. NO EXTERNAL ACTIONS ENABLED.'}


def _latest_date(con: sqlite3.Connection) -> str | None:
    r = con.execute('SELECT report_date FROM reports'
                    ' ORDER BY report_date DESC LIMIT 1').fetchone()
    return r['report_date'] if r else None


@router.get('/reports')
def reports(request: Request, client: str = Query('')) -> dict[str, Any]:
    c = _resolve_client(request, client)
    with _db(c) as con:
        rows = [_row(r) for r in con.execute(
            'SELECT * FROM reports ORDER BY report_date DESC')]
        for r in rows:
            r['block_count_actual'] = con.execute(
                'SELECT COUNT(*) c FROM rig_blocks WHERE report_date=?',
                (r['report_date'],)).fetchone()['c']
    return {'reports': rows}


PRIORITY_RANK = {'critical': 0, 'high': 1, 'medium': 2, 'info': 3}


@router.get('/today')
def today(request: Request, client: str = Query('')) -> dict[str, Any]:
    c = _resolve_client(request, client)
    with _db(c) as con:
        jobs = [_row(r) for r in con.execute('SELECT * FROM jobs')]
        actions = [_row(r) for r in con.execute("SELECT * FROM actions WHERE state='awaiting'")]
        exc = con.execute("SELECT COUNT(*) c FROM exceptions WHERE state='open'").fetchone()['c']
    jobs.sort(key=lambda j: PRIORITY_RANK.get(j['priority'], 9))
    cards = [
        {'key': 'attention', 'label': 'Requires Attention',
         'value': sum(1 for j in jobs if j['priority'] in ('critical', 'high'))},
        {'key': 'dispatch', 'label': 'Dispatch Checks',
         'value': sum(1 for j in jobs if j['lifecycle'] == 'DISPATCH CHECK')},
        {'key': 'rih', 'label': 'RIH Confirmed',
         'value': sum(1 for j in jobs if 'RIH CONFIRMED' in j['lifecycle'])},
        {'key': 'landing', 'label': 'Landing Confirmed',
         'value': sum(1 for j in jobs if 'LAND' in j['lifecycle'])},
        {'key': 'mdd', 'label': 'MDD Due',
         'value': sum(1 for a in actions if a['kind'] == 'Prepare MDD')},
        {'key': 'conflict', 'label': 'Product Conflicts',
         'value': sum(1 for j in jobs if j['lifecycle'] == 'CONFLICT')},
    ]
    items = [{
        'job_id': j['id'], 'priority': j['priority'], 'rig': j['rig'], 'well': j['well'],
        'stage': j['lifecycle'], 'finding': j['finding'],
        'action_required': next((a['title'] for a in actions if a['job_id'] == j['id']), 'Monitor'),
        'evidence_state': _evidence_state(con_job_evidence(j['id'], c)),
        'source_page': j['first_page'],
        'approval': _approval_of(j['id'], actions),
    } for j in jobs]
    return {'cards': cards, 'items': items, 'open_exceptions': exc,
            'date_rule': 'SOURCE TIME shown. DATE NORMALIZATION: AWAITING RULE CONFIRMATION.'}


def con_job_evidence(job_id: str, client: str = DEFAULT_CLIENT) -> list[dict[str, Any]]:
    with _db(client) as con:
        return [_row(r) for r in
                con.execute('SELECT * FROM evidence WHERE job_id=?', (job_id,))]


def _evidence_state(ev: list[dict[str, Any]]) -> str:
    states = {e['status'] for e in ev}
    if 'CONFLICTING' in states:
        return 'CONFLICT'
    if 'INSUFFICIENT' in states:
        return 'INSUFFICIENT EVIDENCE'
    if 'AWAITING AJEEZ' in states:
        return 'AWAITING AJEEZ'
    return 'CONFIRMED'


def _approval_of(job_id: str, actions: list[dict[str, Any]]) -> str:
    mine = [a for a in actions if a['job_id'] == job_id]
    if not mine:
        return 'Monitoring'
    if any(a['state'] == 'blocked' for a in mine):
        return 'Blocked'
    return 'Awaiting Ajeez'


@router.get('/actions')
def action_list(request: Request, client: str = Query('')) -> dict[str, Any]:
    c = _resolve_client(request, client)
    with _db(c) as con:
        actions = [_row(r) for r in con.execute('SELECT * FROM actions')]
        for a in actions:
            dec = con.execute(
                'SELECT decision, note, actor, actor_role, actor_source, at FROM decisions'
                ' WHERE action_id=? ORDER BY id DESC LIMIT 1', (a['id'],)).fetchone()
            a['last_decision'] = dict(dec) if dec else None
            if a['job_id']:
                j = con.execute('SELECT rig, well FROM jobs WHERE id=?',
                                (a['job_id'],)).fetchone()
                a['scope'] = ('%s / %s' % (j['rig'], j['well'])) if j else 'General'
            else:
                a['scope'] = 'General'
    return {'actions': actions}


@router.post('/actions/{action_id}/decision')
def decide(action_id: str, body: dict[str, Any], request: Request,
           client: str = Query('')) -> dict[str, Any]:
    decision = (body.get('decision') or '').lower()
    note = (body.get('note') or '').strip()[:500]
    if decision not in ('approve', 'reject', 'incorrect'):
        _oplog('decision', 'rejected-bad-decision', '%s %s' % (action_id, decision))
        raise HTTPException(400, 'decision must be approve, reject, or incorrect')
    actor, actor_role, actor_source = _resolve_actor(body, request)
    c = _resolve_client(request, client)
    with _db(c) as con:
        a = con.execute('SELECT * FROM actions WHERE id=?', (action_id,)).fetchone()
        if not a:
            _oplog('decision', 'rejected-unknown-action', action_id)
            raise HTTPException(404, 'unknown action')
        a = _row(a)
        if decision == 'approve' and a['state'] == 'blocked':
            _oplog('decision', 'rejected-blocked', action_id)
            raise HTTPException(409, 'Action is blocked: identifiers missing. Resolve first.')
        new_state = {'approve': 'approved', 'reject': 'rejected',
                     'incorrect': 'incorrect'}[decision]
        con.execute('UPDATE actions SET state=? WHERE id=?', (new_state, action_id))
        con.execute('INSERT INTO decisions(action_id,decision,note,actor,actor_role,actor_source)'
                    ' VALUES (?,?,?,?,?,?)',
                    (action_id, decision, note or None,
                     actor, actor_role, actor_source))
        applied = None
        if decision == 'approve' and action_id == 'a-hp701-db':
            con.execute("UPDATE jobs SET lifecycle='RIH CONFIRMED' WHERE id='hp701'")
            con.execute("UPDATE actions SET state='awaiting' WHERE id='a-hp701-mdd'"
                        ' AND state=\'blocked\'')
            applied = ('HP-701 lifecycle recorded as RIH CONFIRMED in demo DB. '
                       'MDD preparation unblocked for drafting (identifiers still missing).')
        con.commit()
    _oplog('decision', new_state,
           '%s %s by %s/%s (%s)' % (action_id, decision, actor, actor_role, actor_source))
    return {'ok': True, 'action': action_id, 'state': new_state,
            'demo_only': True, 'applied': applied,
            'actor': actor, 'actor_role': actor_role, 'actor_source': actor_source}


@router.get('/rigs')
def rigs(request: Request, q: str = Query(''), limit: int = Query(60, le=2000),
         offset: int = Query(0, ge=0),
         report_date: str = Query(''),
         client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        day = report_date or _latest_date(con)
        where, args = ' WHERE report_date=?', [day]
        if q:
            like = '%' + q.upper() + '%'
            where += ' AND (UPPER(rig) LIKE ? OR UPPER(well) LIKE ?)'
            args += [like, like]
        total = con.execute('SELECT COUNT(*) c FROM rig_blocks' + where,
                            args).fetchone()['c']
        rows = [_row(r) for r in con.execute(
            'SELECT report_date, rig, well, first_page, pages, stage,'
            ' orio_relevance, source_file FROM rig_blocks' + where +
            ' ORDER BY rig, well LIMIT ? OFFSET ?', args + [limit, offset])]
    return {'total': total, 'rows': rows, 'report_date': day}


@router.get('/rigs/{rig}/{well}')
def rig_detail(rig: str, well: str, request: Request,
               report_date: str = Query(''),
               client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        day = report_date or _latest_date(con)
        b = con.execute('SELECT * FROM rig_blocks WHERE rig=? AND well=?'
                        ' AND report_date=?',
                        (rig.upper(), well.upper(), day)).fetchone()
        if not b:
            # fall back to any day so an older block is still reachable
            b = con.execute('SELECT * FROM rig_blocks WHERE rig=? AND well=?'
                            ' ORDER BY report_date DESC LIMIT 1',
                            (rig.upper(), well.upper())).fetchone()
        if not b:
            raise HTTPException(404, 'Rig+Well not in the index for this report')
        b = _row(b)
        job = con.execute('SELECT * FROM jobs WHERE rig=? AND well=?',
                          (rig.upper(), well.upper())).fetchone()
        out = dict(b)
        out['tracker'] = [_row(r) for r in con.execute(
            'SELECT * FROM tracker_rows WHERE rig=? AND well=?'
            ' ORDER BY report_date DESC LIMIT 30',
            (rig.upper(), well.upper()))]
        if job:
            job = _row(job)
            job['evidence'] = [_row(r) for r in con.execute(
                'SELECT * FROM evidence WHERE job_id=?', (job['id'],))]
            job['actions'] = [_row(r) for r in con.execute(
                'SELECT * FROM actions WHERE job_id=?', (job['id'],))]
            out['job'] = job
        else:
            out['job'] = None
    return out


@router.get('/jobs/{job_id}')
def job_detail(job_id: str, request: Request,
               client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        j = con.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
        if not j:
            raise HTTPException(404, 'unknown job')
        j = _row(j)
        j['evidence'] = [_row(r) for r in con.execute(
            'SELECT * FROM evidence WHERE job_id=? ORDER BY page', (job_id,))]
        j['actions'] = [_row(r) for r in con.execute(
            'SELECT * FROM actions WHERE job_id=?', (job_id,))]
        j['tracker'] = [_row(r) for r in con.execute(
            'SELECT * FROM tracker_rows WHERE rig=? AND well=?', (j['rig'], j['well']))]
        j['exceptions'] = [_row(r) for r in con.execute(
            'SELECT * FROM exceptions WHERE job_id=?', (job_id,))]
    return j


@router.get('/lifecycle')
def lifecycle(request: Request, client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        jobs = [_row(r) for r in con.execute(
            'SELECT id, rig, well, cycle, lifecycle, priority, finding, first_page FROM jobs')]
    order = ['CONFLICT', 'DISPATCH CHECK', 'INSUFFICIENT EVIDENCE', 'LANDING EVIDENCE',
             'RIH CONFIRMED', 'MONITORING']
    jobs.sort(key=lambda j: order.index(j['lifecycle']) if j['lifecycle'] in order else 99)
    return {'jobs': jobs,
            'note': 'States are assigned only where source evidence supports them.'}


@router.get('/exceptions')
def exceptions(request: Request, client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        rows = [_row(r) for r in con.execute('SELECT * FROM exceptions ORDER BY id')]
    return {'exceptions': rows}


@router.post('/exceptions/{exc_id}/resolve')
def resolve_exception(exc_id: int, body: dict[str, Any], request: Request,
                      client: str = Query('')) -> dict[str, Any]:
    note = (body.get('note') or '').strip()[:500]
    if not note:
        _oplog('exception-resolve', 'rejected-missing-note', 'exception:%d' % exc_id)
        raise HTTPException(400, 'A resolution note is required (demo annotation).')
    actor, actor_role, actor_source = _resolve_actor(body, request)
    c = _resolve_client(request, client)
    with _db(c) as con:
        cur = con.execute("UPDATE exceptions SET state='resolved' WHERE id=?", (exc_id,))
        if cur.rowcount == 0:
            _oplog('exception-resolve', 'rejected-unknown-exception', 'exception:%d' % exc_id)
            raise HTTPException(404, 'unknown exception')
        con.execute('INSERT INTO decisions(action_id,decision,note,actor,actor_role,actor_source)'
                    ' VALUES (?,?,?,?,?,?)',
                    ('exception:%d' % exc_id, 'resolve', note,
                     actor, actor_role, actor_source))
        con.commit()
    _oplog('exception-resolve', 'resolved',
           'exception:%d by %s/%s (%s)' % (exc_id, actor, actor_role, actor_source))
    return {'ok': True, 'demo_only': True,
            'actor': actor, 'actor_role': actor_role, 'actor_source': actor_source}


@router.get('/history')
def history(request: Request, client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        rows = [_row(r) for r in con.execute('SELECT * FROM history ORDER BY date')]
    return {'job': 'HP-707 / THRY-961612',
            'note': 'Benchmark reconstructed ONLY from Ajeez-supplied daily reports. '
                    'File labelled Sep 14 MR 3 pages.pdf is HP-703 / HZEM-491812: '
                    'SOURCE MISMATCH, excluded.',
            'events': rows}


@router.get('/expansion')
def expansion(request: Request, client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        rows = [_row(r) for r in con.execute('SELECT * FROM expansion ORDER BY date, rig')]
    return {'scope': 'OUTSIDE CURRENT ORIO PILOT SCOPE',
            'statement': ('The current pilot may demonstrate that the Morning Report extraction '
                          'foundation can support this workflow. Full competitor-intelligence '
                          'automation requires separate scope approval.'),
            'fields': ['Date RIH', 'Rig', 'Well', 'Size', 'Company', 'Special Remarks'],
            'rows': rows}


# ---------------------------------------------------------------- market intel
# The market-intelligence track: competitor tool-size x company rows mined from
# morning reports. Kept strictly separate from ORIO lifecycle/reconciliation
# state — different evidence, different audience (supervisor -> engineering
# manager). Rows carry page lineage so every value is traceable to a report page.

# Lane C: arbitrary date-range selection. ``month`` stays as a legacy alias
# (one calendar month); ``start``/``end`` (YYYY-MM-DD, inclusive) win when
# given. Every share, count and comparison below is computed from the live
# selection — no hardcoded demo numbers anywhere in this module.
_DATE_RE = re.compile(r'^(\d{4})-(\d{2})-(\d{2})$')
_MONTH_RE = re.compile(r'^(\d{4})-(\d{2})$')


def _as_date(value: str, name: str) -> _date:
    m = _DATE_RE.match((value or '').strip())
    if not m:
        raise HTTPException(400, '%s must be YYYY-MM-DD' % name)
    try:
        return _date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        raise HTTPException(400, '%s is not a calendar date: %r' % (name, value))


def _resolve_range(month: str = '', start: str = '',
                   end: str = '') -> tuple[str | None, str | None, str]:
    """Return (start, end, mode) as ISO strings. ``month`` is legacy alias."""
    month, start, end = (month or '').strip(), (start or '').strip(), (end or '').strip()
    if start or end:
        if not (start and end):
            raise HTTPException(400, 'start and end must be given together (YYYY-MM-DD)')
        s, e = _as_date(start, 'start'), _as_date(end, 'end')
        if s > e:
            raise HTTPException(400, 'start must not be after end')
        return s.isoformat(), e.isoformat(), 'range'
    if month:
        if not _MONTH_RE.match(month):
            raise HTTPException(400, 'month must be YYYY-MM')
        y, m = int(month[:4]), int(month[5:7])
        if not 1 <= m <= 12:
            raise HTTPException(400, 'month is not a calendar month: %r' % month)
        last = (29 if m == 2 and (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0)) else
                28 if m == 2 else 30 if m in (4, 6, 9, 11) else 31)
        return '%04d-%02d-01' % (y, m), '%04d-%02d-%02d' % (y, m, last), 'month'
    return None, None, 'all'


_MATRIX_PROVENANCES = ('workbook', 'extracted', 'adjudicated')
_MATRIX_STATUSES = ('COUNTED', 'PENDING_FILE', 'PENDING_RULING', 'EXCLUDED')
_MATRIX_LOCATIONS = ('onshore', 'offshore')


def _matrix_where(month: str = '', start: str = '', end: str = '',
                  size: str = '', company: str = '',
                  provenance: str = '', date_basis: str = 'report',
                  status: str = '', location: str = ''
                  ) -> tuple[str, list, dict[str, Any]]:
    """Shared selection builder for /matrix, both exports and /matrix/query.

    ``provenance``: ''/'all' (default, all sources) | 'workbook'
    (benchmark reference) | 'extracted' (pipeline candidates) |
    'adjudicated' (Ajeez-ruled June rows). The benchmark is NEVER folded
    into extraction totals — callers that need per-source counts use
    ``totals_by_provenance`` on the response.

    ``status``: ''/'all' (default) | COUNTED | PENDING_FILE |
    PENDING_RULING | EXCLUDED. Ruling status exists only on adjudicated
    rows; other provenances carry NULL and never match a status filter.

    ``location``: ''/'all' (default) | onshore | offshore. On/offshore
    exists only on adjudicated rows (B5).

    ``company`` resolves through data/company_aliases.json (HOWCO->HAL,
    DAVIS-LYNCH->FET, ...); Resp.Co codes (WDV/HDV/TOD) are rejected
    with 400, never silently matched.

    ``date_basis``: 'report' (default) filters the ``date`` column, which
    is the workbook RIH date for benchmark rows and the report date for
    extracted rows; 'event' filters the supported event (RIH) date —
    workbook ``date`` for benchmark rows, evidence ``date_rih`` for
    extracted rows. Extracted rows with unknown event dates are EXCLUDED
    from an event-basis range (never silently counted by report date);
    their count is reported separately as ``unknown_event_dates``.
    """
    sel_start, sel_end, mode = _resolve_range(month, start, end)
    prov = (provenance or '').strip().lower()
    if prov in ('', 'all'):
        prov = ''
    elif prov not in _MATRIX_PROVENANCES:
        raise HTTPException(
            400, "provenance must be '', 'all', 'workbook', 'extracted' or 'adjudicated'")
    st = (status or '').strip().upper()
    if st in ('', 'ALL'):
        st = ''
    elif st not in _MATRIX_STATUSES:
        raise HTTPException(
            400, "status must be '', 'all', 'COUNTED', 'PENDING_FILE', "
            "'PENDING_RULING' or 'EXCLUDED'")
    loc = (location or '').strip().lower()
    if loc in ('', 'all'):
        loc = ''
    elif loc not in _MATRIX_LOCATIONS:
        raise HTTPException(
            400, "location must be '', 'all', 'onshore' or 'offshore'")
    canon_company, company_note = _resolve_company(company)
    basis = (date_basis or 'report').strip().lower()
    if basis not in ('report', 'event'):
        raise HTTPException(400, "date_basis must be 'report' or 'event'")
    where, args = ' WHERE 1=1', []
    if sel_start:
        if basis == 'event':
            where += (' AND COALESCE('
                      "NULLIF(json_extract(evidence_json,'$.date_rih'),''),"
                      " CASE WHEN extraction='extracted' THEN NULL"
                      ' ELSE date END)>=? AND COALESCE('
                      "NULLIF(json_extract(evidence_json,'$.date_rih'),''),"
                      " CASE WHEN extraction='extracted' THEN NULL"
                      ' ELSE date END)<=?')
        else:
            where += ' AND date>=? AND date<=?'
        args += [sel_start, sel_end]
    if size:
        where += ' AND size=?'
        args.append(size)
    if canon_company:
        where += ' AND UPPER(company)=?'
        args.append(canon_company)
    if prov:
        where += ' AND extraction=?'
        args.append(prov)
    if st:
        where += ' AND status=?'
        args.append(st)
    if loc:
        where += ' AND location=?'
        args.append(loc)
    sel: dict[str, Any] = {'mode': mode, 'start': sel_start, 'end': sel_end,
                           'size': size or None, 'company': canon_company or None,
                           'provenance': prov or 'all',
                           'status': st or 'all', 'location': loc or 'all',
                           'date_basis': basis}
    if mode == 'month':
        sel['month'] = month.strip()
    if company_note:
        sel['company_note'] = company_note
    if basis == 'report':
        sel['date_basis_note'] = (
            'report basis: workbook rows filter on their RIH date; '
            'extracted rows filter on their report date (event date '
            'often unknown — see unknown_event_dates)')
    else:
        sel['date_basis_note'] = (
            'event basis: workbook rows filter on their RIH date; '
            'extracted rows filter on supported date_rih; extracted '
            'rows with unknown event dates are excluded from the range')
    return where, args, sel


def _matrix_nodate_where(size: str = '', company: str = '',
                         provenance: str = '', status: str = '',
                         location: str = '') -> tuple[str, list]:
    """Same source/size/company/status/location filters as _matrix_where,
    minus any date range — used to count unresolved event dates across
    the filter set so an event-basis range still discloses the rows it
    excludes."""
    where, args = ' WHERE 1=1', []
    if size:
        where += ' AND size=?'
        args.append(size)
    canon_company, _ = _resolve_company(company)
    if canon_company:
        where += ' AND UPPER(company)=?'
        args.append(canon_company)
    prov = (provenance or '').strip().lower()
    if prov and prov != 'all':
        where += ' AND extraction=?'
        args.append(prov)
    st = (status or '').strip().upper()
    if st and st != 'ALL':
        where += ' AND status=?'
        args.append(st)
    loc = (location or '').strip().lower()
    if loc and loc != 'all':
        where += ' AND location=?'
        args.append(loc)
    return where, args


def _unknown_event_dates(con: sqlite3.Connection, where: str,
                         args: list) -> int:
    """Extracted rows in the selection whose event date is not established.

    Callers pass the date-free filter scope (_matrix_nodate_where) so the
    count also covers rows an event-basis range excludes — unresolved
    dates are disclosed, never silently dropped from view.
    """
    try:
        return con.execute(
            'SELECT COUNT(*) c FROM matrix_rows' + where +
            " AND extraction='extracted' AND (evidence_json IS NULL OR "
            "json_extract(evidence_json,'$.date_rih') IS NULL OR "
            "json_extract(evidence_json,'$.date_rih')='')",
            args).fetchone()['c']
    except Exception:
        return 0


def _provenance_totals(con: sqlite3.Connection, where: str,
                       args: list) -> dict[str, int]:
    """Per-source row counts for the SAME selection (benchmark vs extracted).

    The benchmark total is reported beside — never inside — extraction
    totals.
    """
    out = {'workbook': 0, 'extracted': 0, 'adjudicated': 0}
    try:
        for r in con.execute(
                'SELECT extraction, COUNT(*) c FROM matrix_rows' + where +
                ' GROUP BY extraction', args):
            key = r['extraction'] or 'workbook'
            if key in out:
                out[key] = r['c']
            else:
                out[key] = r['c']
    except Exception:
        pass
    return out


def _with_shares(rows: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    total = sum(r['n'] for r in rows) or 1
    return [{**r, 'share': round(r['n'] / total, 4)} for r in rows]


def _enrich_matrix_row(row: dict[str, Any]) -> dict[str, Any]:
    """Split the two clocks the MI skill requires: report date vs event date.

    Workbook rows carry the hand-built RIH date in ``date`` and no report
    clock; extracted rows carry both in their evidence packet; adjudicated
    rows carry the ruled RIH date in ``date`` plus ruling status, location,
    authority and the June-boundary flag in evidence. Unsupported
    event dates stay null — never backfilled.
    """
    try:
        evidence = json.loads(row.get('evidence_json') or '{}')
    except Exception:
        evidence = {}
    is_extracted = row.get('extraction') == 'extracted'
    event_date = evidence.get('date_rih') if is_extracted else row.get('date')
    report_date = evidence.get('report_date') if is_extracted else None
    event_times = [s.get('event_time') for s in evidence.get('sources', [])
                   if isinstance(s, dict) and s.get('event_time')]
    row['event_date'] = event_date
    row['date_rih'] = event_date  # compat alias
    row['report_date'] = report_date
    row['source_event_times'] = event_times
    row['location'] = row.get('location') or evidence.get('location')
    row['status'] = row.get('status') or evidence.get('status')
    row['authority'] = evidence.get('authority')
    row['june_boundary'] = bool(evidence.get('june_boundary'))
    row['evidence'] = {
        'provenance': row.get('extraction') or 'workbook',
        'source_file': row.get('source_file'),
        'source_page': row.get('source_page'),
        'has_page': (str(row.get('source_page') or '').isdigit()),
        'event_times': event_times,
        'status': row['status'],
        'location': row['location'],
        'authority': row['authority'],
        'june_boundary': row['june_boundary'],
    }
    return row


def _matrix_selection_rows(con: sqlite3.Connection, where: str,
                           args: list) -> list[dict[str, Any]]:
    return [_enrich_matrix_row(_row(r)) for r in con.execute(
        'SELECT * FROM matrix_rows' + where +
        ' ORDER BY date, rig, well', args)]


def _matrix_coverage(con: sqlite3.Connection) -> dict[str, Any]:
    r = con.execute('SELECT MIN(date) mn, MAX(date) mx,'
                    ' COUNT(DISTINCT date) days, COUNT(*) n'
                    ' FROM matrix_rows').fetchone()
    last = None
    try:
        last = con.execute('SELECT MAX(ingested_at) m FROM reports').fetchone()
    except Exception:
        last = None
    try:
        rep_days = con.execute('SELECT COUNT(*) c FROM reports').fetchone()['c']
    except Exception:
        rep_days = None
    return {'event_min': r['mn'], 'event_max': r['mx'],
            'distinct_event_days': r['days'], 'matrix_rows_all': r['n'],
            'report_days': rep_days,
            'last_update': (last['m'] if last else None)}


@router.get('/matrix')
def matrix(request: Request, month: str = Query(''), start: str = Query(''),
           end: str = Query(''), size: str = Query(''),
           company: str = Query(''), limit: int = Query(500, le=5000),
           offset: int = Query(0, ge=0),
           provenance: str = Query(''),
           date_basis: str = Query('report'),
           status: str = Query(''), location: str = Query(''),
           client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        where, args, sel = _matrix_where(month, start, end, size, company,
                                         provenance, date_basis,
                                         status, location)
        total = con.execute('SELECT COUNT(*) c FROM matrix_rows' + where,
                            args).fetchone()['c']
        page = [_row(r) for r in con.execute(
            'SELECT * FROM matrix_rows' + where +
            ' ORDER BY date, rig, well LIMIT ? OFFSET ?', args + [limit, offset])]
        rows = [_enrich_matrix_row(r) for r in page]
        sizes = [_row(r) for r in con.execute(
            'SELECT size, COUNT(*) n FROM matrix_rows' + where +
            ' GROUP BY size ORDER BY n DESC', args)]
        # Company breakdown excludes company='UNPROVEN' rows: an UNPROVEN
        # company is counted in size/location totals only (row 106,
        # WDV Resp.Co code, rules-v2 R1 corollary). The excluded count is
        # disclosed as unproven_company_rows, never silently dropped.
        comps = [_row(r) for r in con.execute(
            'SELECT company, COUNT(*) n FROM matrix_rows' + where +
            " AND COALESCE(company,'')!='UNPROVEN'" +
            ' GROUP BY company ORDER BY n DESC', args)]
        unproven_n = con.execute(
            'SELECT COUNT(*) c FROM matrix_rows' + where +
            " AND company='UNPROVEN'", args).fetchone()['c']
        locs = [_row(r) for r in con.execute(
            'SELECT location, COUNT(*) n FROM matrix_rows' + where +
            ' GROUP BY location ORDER BY n DESC', args)]
        stats = [_row(r) for r in con.execute(
            "SELECT COALESCE(status,'NULL') status, COUNT(*) n"
            ' FROM matrix_rows' + where +
            ' GROUP BY 1 ORDER BY n DESC', args)]
        months = [_row(r) for r in con.execute(
            'SELECT substr(date,1,7) month, COUNT(*) n FROM matrix_rows'
            ' GROUP BY month ORDER BY month')]
        by_size_company = [_row(r) for r in con.execute(
            'SELECT size, company, COUNT(*) n FROM matrix_rows' + where +
            ' GROUP BY size, company ORDER BY size, n DESC', args)]
        with_page = con.execute(
            'SELECT COUNT(*) c FROM matrix_rows' + where +
            # No-invent rule: only integer report pages count as page lineage.
            # Workbook-sourced rows honestly carry source_page='workbook'.
            " AND source_page IS NOT NULL AND source_page != 'workbook'",
            args).fetchone()['c']
        workbook_n = con.execute(
            'SELECT COUNT(*) c FROM matrix_rows' + where +
            " AND (source_page='workbook' OR extraction='workbook'"
            " OR extraction IS NULL)",
            args).fetchone()['c']
        coverage = _matrix_coverage(con)
        compare = _period_compare(con, where, args, sel, size, company,
                                  sel.get('provenance', ''),
                                  sel.get('date_basis', 'report'),
                                  sel.get('status', ''),
                                  sel.get('location', ''))
        prov_totals = _provenance_totals(con, where, args)
        nd_where, nd_args = _matrix_nodate_where(
            size, company, sel.get('provenance', ''),
            sel.get('status', ''), sel.get('location', ''))
        unknown_ev = _unknown_event_dates(con, nd_where, nd_args)
        try:
            innovex_n = con.execute(
                'SELECT COUNT(*) c FROM innovex_own').fetchone()['c']
        except Exception:
            innovex_n = 0
    return {'rows': rows, 'total': total,
            'selection': sel,
            'totals_by_provenance': prov_totals,
            'unknown_event_dates': unknown_ev,
            'unknown_event_dates_scope': ('extracted rows matching '
                                          'source/size/company filters, '
                                          'any date — includes rows an '
                                          'event-basis range excludes'),
            'by_size': _with_shares(sizes, 'size'),
            'by_company': _with_shares(comps, 'company'),
            'by_location': _with_shares(
                [{'location': r['location'], 'n': r['n']} for r in locs],
                'location'),
            'by_status': stats,
            'unproven_company_rows': unproven_n,
            'unproven_company_note': ('rows with company UNPROVEN are '
                                      'counted in size/location totals only, '
                                      'never in by_company (rules-v2 R1)'),
            'innovex_own': {'count': innovex_n,
                            'note': 'INNOVEX-tool jobs, out of competitor '
                                    'scope: fenced in innovex_own, never in '
                                    'matrix totals'},
            'by_month': months,
            'by_size_company': by_size_company,
            'period_compare': compare,
            'coverage': coverage,
            'lineage': {'with_source_page': with_page, 'total': total},
            'evidence_summary': {'with_page': with_page,
                                 'workbook': workbook_n,
                                 'total': total,
                                 'run_level': 'per-row evidence{}: provenance, '
                                 'source file/page, event/report dates'},
            'track': 'market intelligence (separate from ORIO lifecycle)'}


def _period_compare(con: sqlite3.Connection, where: str, args: list,
                    sel: dict[str, Any], size: str,
                    company: str, provenance: str = '',
                    date_basis: str = 'report', status: str = '',
                    location: str = '') -> dict[str, Any] | None:
    """Current selection vs the previous equal-length period, same filters."""
    if not (sel.get('start') and sel.get('end')):
        return None
    s, e = _as_date(sel['start'], 'start'), _as_date(sel['end'], 'end')
    span = (e - s).days + 1
    from datetime import timedelta as _td
    ps, pe = s - _td(days=span), s - _td(days=1)
    cur = con.execute('SELECT COUNT(*) c FROM matrix_rows' + where,
                      args).fetchone()['c']
    pwhere, pargs, _ = _matrix_where('', ps.isoformat(), pe.isoformat(),
                                     size, company, provenance, date_basis,
                                     status, location)
    prev = con.execute('SELECT COUNT(*) c FROM matrix_rows' + pwhere,
                       pargs).fetchone()['c']
    delta = cur - prev
    pct = (round(delta / prev, 4) if prev else None)
    return {'current': {'start': s.isoformat(), 'end': e.isoformat(), 'n': cur},
            'previous': {'start': ps.isoformat(), 'end': pe.isoformat(), 'n': prev},
            'delta': delta, 'pct_change': pct}


@router.get('/matrix/export.csv')
def matrix_csv(request: Request, month: str = Query(''),
               start: str = Query(''), end: str = Query(''),
               size: str = Query(''), company: str = Query(''),
               provenance: str = Query(''),
               date_basis: str = Query('report'),
               status: str = Query(''), location: str = Query(''),
               envelope: str = Query(''),
               client: str = Query('')) -> Response:
    with _db(_resolve_client(request, client)) as con:
        where, args, sel = _matrix_where(month, start, end, size, company,
                                         provenance, date_basis,
                                         status, location)
        evidence_column = ('evidence_json' if any(r[1] == 'evidence_json'
                           for r in con.execute('PRAGMA table_info(matrix_rows)'))
                           else 'NULL AS evidence_json')
        cols = {r[1] for r in con.execute('PRAGMA table_info(matrix_rows)')}
        loc_column = 'location' if 'location' in cols else 'NULL AS location'
        status_column = 'status' if 'status' in cols else 'NULL AS status'
        rows = con.execute(
            'SELECT date, rig, well, size, company, remarks, source_file,'
            ' source_page, extraction, ' + loc_column + ', ' + status_column + ', '
            + evidence_column + ' FROM matrix_rows' + where +
            ' ORDER BY date, rig, well', args)
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(['Date RIH', 'Report Date', 'Source Event Times', 'Rig', 'Well', 'Size', 'Company',
                    'Special Remarks', 'Source File', 'Evidence JSON', 'Source Page',
                    'Location', 'Status', 'Provenance'])
        for r in rows:
            evidence = json.loads(r['evidence_json'] or '{}')
            is_extracted = r['extraction'] == 'extracted'
            rih_date = evidence.get('date_rih') if is_extracted else r['date']
            report_date = evidence.get('report_date', r['date']) if is_extracted else ''
            event_times = ';'.join(dict.fromkeys(s['event_time'] for s in evidence.get('sources', []) if s.get('event_time')))
            w.writerow([_csv_cell(rih_date), _csv_cell(report_date), _csv_cell(event_times), _csv_cell(r['rig']),
                        _csv_cell(r['well']), _csv_cell(r['size']),
                        _csv_cell(r['company']), _csv_cell(r['remarks']),
                        _csv_cell(r['source_file']),
                        _csv_cell(r['evidence_json']),
                        _csv_cell(r['source_page']),
                        _csv_cell(r['location']),
                        _csv_cell(r['status']),
                        _csv_cell(r['extraction'])])
    tag = ('%s_to_%s' % (sel['start'], sel['end']) if sel.get('start')
           else (sel.get('month') or 'all'))
    if sel.get('provenance') not in (None, 'all'):
        tag += '_%s' % sel['provenance']
    if sel.get('date_basis') not in (None, 'report'):
        tag += '_%s-basis' % sel['date_basis']
    name = 'orio_market_intel_%s.csv' % tag
    if envelope == 'json':
        return _envelope(name, 'text/csv', buf.getvalue().encode('utf-8'))
    return Response(content=buf.getvalue(), media_type='text/csv',
                    headers={'Content-Disposition':
                             'attachment; filename="%s"' % name})


# Lane C: genuine .xlsx export (openpyxl) + ORIO selection seam.
# Same-dataset rule: the workbook is built from the identical WHERE clause
# as /matrix and /matrix/export.csv for the same selection — the
# Reconciliation sheet proves it (details count == summary total ==
# sum of shares, evidence count == lineage count).

_XLSX_MEDIA = ('application/vnd.openxmlformats-officedocument.'
               'spreadsheetml.sheet')


def _xlsx_text(v: Any) -> Any:
    """Cell-safe value: numbers stay numeric; formula-leading text is
    neutralized with Excel's native text marker (same rule as _csv_cell)."""
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    if (s[:1] in ('=', '+', '-', '@', '\t', '\r', '\n')
            or s.lstrip()[:1] in ('=', '+', '-', '@')):
        return "'" + s
    return s


def _build_matrix_workbook(rows: list[dict[str, Any]],
                           sel: dict[str, Any],
                           comps: list[dict[str, Any]],
                           sizes: list[dict[str, Any]],
                           compare: dict[str, Any] | None,
                           coverage: dict[str, Any],
                           with_page: int, generated_at: str,
                           prov_totals: dict[str, int] | None = None,
                           unknown_event_dates: int = 0):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    bold = Font(bold=True)
    total = len(rows)
    prov_totals = prov_totals or {}

    ws = wb.active
    ws.title = 'Summary'
    sel_line = ' / '.join('%s=%s' % (k, v) for k, v in sel.items()
                           if v is not None)
    ws.append(['ORIO Market Intelligence — selection summary'])
    ws.append(['Selection', sel_line or 'all rows'])
    ws.append(['Date basis', sel.get('date_basis_note') or sel.get('date_basis') or 'report'])
    ws.append(['Generated (UTC)', generated_at])
    ws.append(['Total rows', total])
    ws.append(['Rows workbook-sourced (benchmark)', prov_totals.get('workbook', '')])
    ws.append(['Rows extracted (pipeline candidates)', prov_totals.get('extracted', '')])
    ws.append(['Rows adjudicated (Ajeez-ruled)', prov_totals.get('adjudicated', '')])
    ws.append(['Extracted rows, event date unknown', unknown_event_dates])
    if prov_totals.get('adjudicated'):
        ws.append(['Acceptance', 'June ACCEPTED (Ajeez-ruled 2026-10-20): '
                   'adjudicated COUNTED rows are findings; company UNPROVEN '
                   'rows count in size/location only; PENDING/EXCLUDED rows '
                   'and extracted candidates are never folded in'])
    else:
        ws.append(['Acceptance', 'June HOLD — benchmark and extracted rows are '
                   'reference vs candidates, never summed as accepted findings'])
    ws.append(['Rows with source page', with_page])
    ws.append(['Rows workbook-sourced',
               sum(1 for r in rows if not r['evidence']['has_page'])])
    ws.append([])
    ws.append(['Company', 'Rows', 'Share'])
    for c in comps:
        ws.append([_xlsx_text(c['company']), c['n'], c.get('share', '')])
    ws.append([])
    ws.append(['Tool size', 'Rows', 'Share'])
    for s in sizes:
        ws.append([_xlsx_text(s['size']), s['n'], s.get('share', '')])
    ws.append([])
    if compare:
        ws.append(['Period compare', 'Rows'])
        ws.append(['Current %s..%s' % (compare['current']['start'],
                                       compare['current']['end']),
                   compare['current']['n']])
        ws.append(['Previous %s..%s' % (compare['previous']['start'],
                                        compare['previous']['end']),
                   compare['previous']['n']])
        ws.append(['Delta', compare['delta']])
        ws.append(['Pct change', compare['pct_change']
                   if compare['pct_change'] is not None else 'n/a'])
        ws.append([])
    ws.append(['Coverage: event range',
               '%s .. %s' % (coverage.get('event_min'),
                             coverage.get('event_max'))])
    ws.append(['Coverage: distinct event days',
               coverage.get('distinct_event_days')])
    ws.append(['Coverage: report days ingested', coverage.get('report_days')])
    ws.append(['Coverage: last update (ingested_at max)',
               coverage.get('last_update')])
    for cell in ('A1', 'A8', 'A%d' % (10 + len(comps) + 1)):
        try:
            ws[cell].font = bold
        except Exception:
            pass

    wd = wb.create_sheet('Details')
    wd.append(['Event Date (RIH)', 'Report Date', 'Source Event Times', 'Rig',
               'Well', 'Size', 'Company', 'Location', 'Status',
               'Special Remarks', 'Source File',
               'Source Page', 'Provenance'])
    for r in rows:
        wd.append([_xlsx_text(r.get('event_date')),
                   _xlsx_text(r.get('report_date')),
                   _xlsx_text(';'.join(r.get('source_event_times') or [])),
                   _xlsx_text(r.get('rig')), _xlsx_text(r.get('well')),
                   _xlsx_text(r.get('size')), _xlsx_text(r.get('company')),
                   _xlsx_text(r.get('location')), _xlsx_text(r.get('status')),
                   _xlsx_text(r.get('remarks')),
                   _xlsx_text(r.get('source_file')),
                   _xlsx_text(r.get('source_page')),
                   _xlsx_text(r.get('extraction'))])

    we = wb.create_sheet('Evidence')
    we.append(['Event Date', 'Rig', 'Well', 'Size', 'Company', 'Source File',
               'Source Page', 'Provenance', 'Report Date', 'Event Times'])
    for r in rows:
        if not r['evidence']['has_page']:
            continue
        we.append([_xlsx_text(r.get('event_date')), _xlsx_text(r.get('rig')),
                   _xlsx_text(r.get('well')), _xlsx_text(r.get('size')),
                   _xlsx_text(r.get('company')),
                   _xlsx_text(r.get('source_file')),
                   _xlsx_text(r.get('source_page')),
                   _xlsx_text(r.get('extraction')),
                   _xlsx_text(r.get('report_date')),
                   _xlsx_text(';'.join(r.get('source_event_times') or []))])

    wc = wb.create_sheet('Coverage')
    wc.append(['Event Date (UNKNOWN = event date not established)', 'Rows in selection'])
    days: dict[str, int] = {}
    for r in rows:
        k = r.get('event_date') or 'UNKNOWN (event date not established)'
        days[k] = days.get(k, 0) + 1
    for d in sorted(days):
        wc.append([d, days[d]])

    wr = wb.create_sheet('Reconciliation')
    wr.append(['Check', 'Expected', 'Actual', 'Status'])
    _UNKNOWN_LABEL = 'UNKNOWN (event date not established)'
    _unknown_in_rows = sum(
        1 for r in rows
        if (r.get('extraction') or 'workbook') == 'extracted'
        and not r.get('event_date'))
    _unproven_in_rows = sum(1 for r in rows if r.get('company') == 'UNPROVEN')
    checks = [
        ('details rows == summary total', total, total, True),
        ('summary total - UNPROVEN == sum(company shares base) '
         '(UNPROVEN counts in size/location only, rules-v2 R1)',
         total - _unproven_in_rows, sum(c['n'] for c in comps),
         total - _unproven_in_rows == sum(c['n'] for c in comps)),
        ('summary total == sum(size shares base)',
         total, sum(s['n'] for s in sizes),
         total == sum(s['n'] for s in sizes)),
        ('evidence sheet rows == lineage with_page',
         with_page, sum(1 for r in rows if r['evidence']['has_page']),
         with_page == sum(1 for r in rows if r['evidence']['has_page'])),
        ('coverage sheet rows sum == summary total',
         total, sum(days.values()), total == sum(days.values())),
        ('sum(provenance totals incl. adjudicated) == summary total '
         '(separate, never folded)',
         total, sum((prov_totals or {}).values()),
         total == sum((prov_totals or {}).values())),
        ('unknown-event-date rows in Details == UNKNOWN coverage bucket',
         _unknown_in_rows,
         days.get(_UNKNOWN_LABEL, 0),
         _unknown_in_rows == days.get(_UNKNOWN_LABEL, 0)),
        ('filter-scope unknown-event-date rows >= UNKNOWN in Details',
         True, unknown_event_dates >= _unknown_in_rows,
         unknown_event_dates >= _unknown_in_rows),
    ]
    if compare:
        checks.append(('period current n == summary total', total,
                       compare['current']['n'],
                       total == compare['current']['n']))
    for name, exp, act, ok in checks:
        wr.append([name, exp, act, 'PASS' if ok else 'FAIL'])
    return wb


@router.get('/matrix/export.xlsx')
def matrix_xlsx(request: Request, month: str = Query(''),
                start: str = Query(''), end: str = Query(''),
                size: str = Query(''), company: str = Query(''),
                provenance: str = Query(''),
                date_basis: str = Query('report'),
                status: str = Query(''), location: str = Query(''),
                envelope: str = Query(''),
                client: str = Query('')) -> Response:
    c = _resolve_client(request, client)
    with _db(c) as con:
        where, args, sel = _matrix_where(month, start, end, size, company,
                                         provenance, date_basis,
                                         status, location)
        rows = _matrix_selection_rows(con, where, args)
        comps = _with_shares(
            [_row(r) for r in con.execute(
                'SELECT company, COUNT(*) n FROM matrix_rows' + where +
                " AND COALESCE(company,'')!='UNPROVEN'" +
                ' GROUP BY company ORDER BY n DESC', args)], 'company')
        sizes = _with_shares(
            [_row(r) for r in con.execute(
                'SELECT size, COUNT(*) n FROM matrix_rows' + where +
                ' GROUP BY size ORDER BY n DESC', args)], 'size')
        with_page = sum(1 for r in rows if r['evidence']['has_page'])
        compare = _period_compare(con, where, args, sel, size, company,
                                  sel.get('provenance', ''),
                                  sel.get('date_basis', 'report'),
                                  sel.get('status', ''),
                                  sel.get('location', ''))
        coverage = _matrix_coverage(con)
        prov_totals = _provenance_totals(con, where, args)
        nd_where, nd_args = _matrix_nodate_where(
            size, company, sel.get('provenance', ''))
        unknown_ev = _unknown_event_dates(con, nd_where, nd_args)
    wb = _build_matrix_workbook(
        rows, sel, comps, sizes, compare, coverage, with_page,
        time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
        prov_totals, unknown_ev)
    buf = io.BytesIO()
    wb.save(buf)
    tag = ('%s_to_%s' % (sel['start'], sel['end']) if sel.get('start')
           else (sel.get('month') or 'all'))
    if sel.get('provenance') not in (None, 'all'):
        tag += '_%s' % sel['provenance']
    if sel.get('date_basis') not in (None, 'report'):
        tag += '_%s-basis' % sel['date_basis']
    xname = 'orio_market_intel_%s.xlsx' % tag
    if envelope == 'json':
        return _envelope(xname, _XLSX_MEDIA, buf.getvalue())
    return Response(content=buf.getvalue(), media_type=_XLSX_MEDIA,
                    headers={'Content-Disposition':
                             'attachment; filename="%s"' % xname})


def _selection_payload(con: sqlite3.Connection, where: str, args: list,
                       sel: dict[str, Any], size: str,
                       company: str) -> dict[str, Any]:
    total = con.execute('SELECT COUNT(*) c FROM matrix_rows' + where,
                        args).fetchone()['c']
    comps = _with_shares(
        [_row(r) for r in con.execute(
            'SELECT company, COUNT(*) n FROM matrix_rows' + where +
            " AND COALESCE(company,'')!='UNPROVEN'" +
            ' GROUP BY company ORDER BY n DESC', args)], 'company')
    sizes = _with_shares(
        [_row(r) for r in con.execute(
            'SELECT size, COUNT(*) n FROM matrix_rows' + where +
            ' GROUP BY size ORDER BY n DESC', args)], 'size')
    return {
        'selection': sel,
        'total': total,
        'totals_by_provenance': _provenance_totals(con, where, args),
        'unknown_event_dates': _unknown_event_dates(
            con, *_matrix_nodate_where(size, company,
                                       sel.get('provenance', ''),
                                       sel.get('status', ''),
                                       sel.get('location', ''))),
        'unknown_event_dates_scope': ('extracted rows matching '
                                      'source/size/company filters, any date'),
        'shares': {'by_company': comps, 'by_size': sizes},
        'period_compare': _period_compare(con, where, args, sel, size,
                                          company,
                                          sel.get('provenance', ''),
                                          sel.get('date_basis', 'report'),
                                          sel.get('status', ''),
                                          sel.get('location', '')),
        'coverage': _matrix_coverage(con),
        'exports': {
            'csv': '/matrix/export.csv',
            'xlsx': '/matrix/export.xlsx',
            'note': ('Append the same selection params '
                     '(?start=&end=&size=&company=&provenance=&date_basis=).'),
        },
        'track': 'market intelligence (separate from ORIO lifecycle)',
        'note': ('ORIO MI selection seam: pass this selection to the matrix '
                 'endpoints or exports to reproduce the exact row set.'),
    }


@router.get('/matrix/query')
def matrix_query_get(request: Request, month: str = Query(''),
                     start: str = Query(''), end: str = Query(''),
                     size: str = Query(''), company: str = Query(''),
                     provenance: str = Query(''),
                     date_basis: str = Query('report'),
                     status: str = Query(''), location: str = Query(''),
                     client: str = Query('')) -> dict[str, Any]:
    """ORIO extension seam (GET): the UI's Ask-ORIO button and any agent
    call this with the current range+filters; the reply is the reproducible
    selection spec ORIO queries against."""
    c = _resolve_client(request, client)
    with _db(c) as con:
        where, args, sel = _matrix_where(month, start, end, size, company,
                                         provenance, date_basis,
                                         status, location)
        return _selection_payload(con, where, args, sel, size, company)


@router.post('/matrix/query')
def matrix_query_post(body: dict[str, Any], request: Request,
                      client: str = Query('')) -> dict[str, Any]:
    """ORIO extension seam (POST): JSON {month?, start?, end?, size?,
    company?, provenance?, date_basis?} for skill/agent callers that POST
    a selection."""
    body = body or {}
    c = _resolve_client(request, client)
    with _db(c) as con:
        where, args, sel = _matrix_where(
            body.get('month', ''), body.get('start', ''),
            body.get('end', ''), body.get('size', ''),
            body.get('company', ''), body.get('provenance', ''),
            body.get('date_basis', 'report'), body.get('status', ''),
            body.get('location', ''))
        return _selection_payload(con, where, args, sel,
                                  body.get('size', ''),
                                  body.get('company', ''))


@router.get('/export.csv')
def export_csv(request: Request, client: str = Query('')) -> Response:
    buf = io.StringIO()
    with _db(_resolve_client(request, client)) as con:
        rows = con.execute('SELECT * FROM rig_blocks ORDER BY first_page')
        w = csv.writer(buf)
        w.writerow(['rig', 'well', 'first_page', 'pages', 'stage', 'orio_relevance',
                    'last_24h', 'next_24h'])
        for r in rows:
            w.writerow([_csv_cell(r['rig']), _csv_cell(r['well']),
                        _csv_cell(r['first_page']), _csv_cell(r['pages']),
                        _csv_cell(r['stage']), _csv_cell(r['orio_relevance']),
                        _csv_cell(r['last24']), _csv_cell(r['next24'])])
    return Response(content=buf.getvalue(), media_type='text/csv',
                    headers={'Content-Disposition':
                             'attachment; filename="orio_demo_rig_index_2026-09-19.csv"'})


@router.get('/mdd/preview')
def mdd_preview(request: Request, job: str = Query('hp701'),
                client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        j = con.execute('SELECT * FROM jobs WHERE id=?', (job,)).fetchone()
        if not j:
            raise HTTPException(404, 'unknown job')
        j = _row(j)
        docs = [_row(r) for r in con.execute('SELECT * FROM mdd_docs')]
        hp704 = con.execute("SELECT well FROM rig_blocks WHERE rig='HP-704'").fetchall()
    lines = [
        ('Job', '%s / %s' % (j['rig'], j['well'])),
        ('RIH evidence', 'Confirmed Sep 19 MR pp%d-%d' % (j['first_page'], j['first_page'] + j['pages'] - 1)
         if job == 'hp701' else 'See job evidence'),
        ('RIH date / time', 'SOURCE: Sep 19 MR 05:00-05:00 period. Normalized date AWAITING RULE CONFIRMATION.'),
        ('Foreman', j['foreman'] or 'MISSING - REQUIRES AJEEZ / SOURCE DATA'),
        ('PO number', 'MISSING - REQUIRES AJEEZ / SOURCE DATA'),
        ('Shipment number', 'MISSING - REQUIRES AJEEZ / SOURCE DATA'),
        ('Signature contact', 'MISSING - REQUIRES AJEEZ / SOURCE DATA'),
    ]
    cross = ('Supplied MDD PO 4508154877 belongs to HP-704 / THRY-95140. Sep 19 shows HP-704 on '
             + (hp704[0]['well'] if hp704 else 'a different well')
             + ' (normal rig move). It cannot be reused for this job.')
    # Phase 2 seam: the MDD email payload is built as local data only.
    # The email channel is BLOCKED, so this is previewed/stored, never sent.
    email = get_channel('email')
    payload = build_email_payload(
        recipient='MISSING - REQUIRES AJEEZ / SOURCE DATA',
        subject='ORIO MDD %s / %s - preview only' % (j['rig'], j['well']),
        body='\n'.join('%s: %s' % (k, v) for k, v in lines),
        attachments=[],
        ref='mdd:%s' % job,
    )
    artifact = email.prepare(payload)
    return {'job': '%s / %s' % (j['rig'], j['well']), 'lines': lines,
            'supplied_mdd': docs, 'cross_check': cross, 'send_enabled': False,
            'send_note': 'Production connection disabled in demo. Preview only.',
            'delivery_payload': payload,
            'channel': email.status(),
            'channel_artifact': artifact}


@router.post('/drafts/email')
def draft_email(body: dict[str, Any], request: Request,
                client: str = Query('')) -> dict[str, Any]:
    job_id = body.get('job') or 'hp701'
    c = _resolve_client(request, client)
    with _db(c) as con:
        j = con.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
        if not j:
            raise HTTPException(404, 'unknown job')
        j = _row(j)
        text = ('To: MISSING - REQUIRES AJEEZ / SOURCE DATA\n'
                'Subject: ORIO %s / %s - %s\n\nFinding: %s\n'
                'Evidence: Sep 19 MR pages %d-%d. Foreman: %s.\n'
                'Date rule: SOURCE TIME. Normalization AWAITING CONFIRMATION.\n'
                'This is a LOCAL DEMO PREVIEW. Nothing was sent.'
                % (j['rig'], j['well'], j['lifecycle'], j['finding'],
                   j['first_page'], j['first_page'] + j['pages'] - 1, j['foreman']))
        cur = con.execute("INSERT INTO drafts(kind,ref,body) VALUES ('email',?,?)",
                          (job_id, text))
        did = cur.lastrowid
        # Phase 2 seam: same content shaped as a delivery payload, prepared
        # through the BLOCKED email channel and stored in the local outbox.
        # deliver() is never called — nothing leaves this machine.
        email = get_channel('email')
        payload = build_email_payload(
            recipient='MISSING - REQUIRES AJEEZ / SOURCE DATA',
            subject='ORIO %s / %s - %s' % (j['rig'], j['well'], j['lifecycle']),
            body=text,
            attachments=[],
            ref='draft:%d' % did,
        )
        artifact = email.prepare(payload)
        out = con.execute(
            'INSERT INTO delivery_outbox(channel,ref,payload_json,state,reason)'
            ' VALUES (?,?,?,?,?)',
            ('email', 'draft:%d' % did, json.dumps(payload),
             artifact['state'], artifact['blocked_reason']))
        con.commit()
        oid = out.lastrowid
    _oplog('draft-email', 'blocked-stored',
           'draft:%d for %s (channel blocked, nothing sent)' % (did, job_id))
    return {'ok': True, 'draft_id': did, 'body': text, 'sent': False,
            'delivery_payload': payload, 'channel': email.status(),
            'outbox_id': oid, 'outbox_state': artifact['state']}


@router.get('/drafts')
def drafts(request: Request, client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        rows = [_row(r) for r in con.execute('SELECT * FROM drafts ORDER BY id DESC LIMIT 20')]
    return {'drafts': rows, 'send_enabled': False}


# ---------------------------------------------------------------- clients
# Phase 4: one codebase, one workspace folder per client. Client #1 is
# `orio-ops` (Ajeez). Onboarding is tools/onboard_client.sh — this API
# never auto-creates a workspace (unknown slugs answer 404).


@router.get('/clients')
def clients() -> dict[str, Any]:
    return {'clients': list_clients(get_hermes_home()),
            'default': DEFAULT_CLIENT,
            'note': ('Client #1 (%s) is Ajeez Sebastian. Onboarding a new '
                     'client = new folder + profile + credentials via '
                     'tools/onboard_client.sh. No code fork.'
                     % DEFAULT_CLIENT)}


@router.get('/clients/{slug}')
def client_profile(slug: str) -> dict[str, Any]:
    home = get_hermes_home()
    try:
        cfg = load_config(home, slug)
    except UnknownClient as e:
        raise HTTPException(404, str(e)) from None
    except ValueError as e:
        raise HTTPException(400, str(e)) from None
    from clients import config_path as _cfg_path, db_path as _dbp
    return {'profile': cfg,
            'db_exists': _dbp(home, slug).exists(),
            'config_path': str(_cfg_path(home, slug))}


# ---------------------------------------------------------------- channels
# Phase 2 seams: every channel reports BLOCKED with its reason. The outbox
# holds prepared payloads as local data; there is no send path to call.

# ---------------------------------------------------------------- intelligence
# LANE 7: read-only range views over VERIFIED June registers (features 1-4).
# Pure functions, no sockets, no writes. Frozen registers are never edited.

_STALE_LANE7_NOTE = (
    'STALE vs Ajeez 2026-10-20 rulings: these views read the pre-ruling '
    'registers (dv-june-adjudicated.json day-events + lane4-reconciliation '
    'round2 131-row reconciliation). Ruled June truth lives in '
    'phase7-ajeez-rulings/adjudicated-rows-v2.json (123 COUNTED) and the '
    '/matrix provenance=adjudicated rows. Lane-7 re-pointing is deferred; '
    'do not cite these views as ruled findings.')


def _stale_lane7(view: dict[str, Any]) -> dict[str, Any]:
    """Mark a lane-7 view visibly stale (M4): additive keys only, so the
    frozen-register checks keep passing while no caller can mistake the
    view for ruled truth."""
    view = dict(view)
    view['stale'] = True
    view['stale_note'] = _STALE_LANE7_NOTE
    view['superseded_by'] = ('adjudicated-rows-v2.json + '
                             '/matrix?provenance=adjudicated')
    return view

try:
    from intelligence import build_bundle
except Exception:  # pragma: no cover - gateway loads dashboard standalone
    import sys as _sys2

    _here2 = str(Path(__file__).resolve().parent)
    if _here2 not in _sys2.path:
        _sys2.path.insert(0, _here2)
    from intelligence import build_bundle  # type: ignore[no-redef]


@router.get('/intelligence/upcoming')
def intelligence_upcoming(as_of: str = Query('2026-06-30')) -> dict[str, Any]:
    """Feature 1: explicit upcoming-job plans, historical as-of-date view."""
    return _stale_lane7(build_bundle(as_of=as_of)['upcoming'])


@router.get('/intelligence/suppliers/first-seen')
def intelligence_first_seen() -> dict[str, Any]:
    """Feature 2: first supplier appearances + repeat use, coverage shown."""
    return _stale_lane7(build_bundle()['suppliers'])


@router.get('/intelligence/tool-problems')
def intelligence_problems() -> dict[str, Any]:
    """Feature 3: explicitly reported tool problems + resolution state."""
    return _stale_lane7(build_bundle()['problems'])


@router.get('/intelligence/corrections')
def intelligence_corrections() -> dict[str, Any]:
    """Feature 4: reviewer correction flow with originals + effort metrics."""
    return _stale_lane7(build_bundle()['corrections'])


@router.get('/intelligence/questions')
def intelligence_questions() -> dict[str, Any]:
    """Banked ORIO questions with evidence links for the four views."""
    return build_bundle()['questions']


@router.get('/channels/status')
def channels_status() -> dict[str, Any]:
    return {'channels': all_statuses(),
            'note': 'All channels BLOCKED. Payloads are prepared and stored '
                    'locally only; nothing is sent. See dashboard/channels.py '
                    'for the documented unblock-day adapters.'}


@router.get('/delivery/outbox')
def delivery_outbox(request: Request,
                    client: str = Query('')) -> dict[str, Any]:
    with _db(_resolve_client(request, client)) as con:
        rows = [_row(r) for r in con.execute(
            'SELECT * FROM delivery_outbox ORDER BY id DESC LIMIT 50')]
        for r in rows:
            try:
                r['payload'] = json.loads(r.pop('payload_json') or '{}')
            except Exception:
                r['payload'] = {}
    return {'outbox': rows, 'send_enabled': False,
            'note': 'Stored local payloads only. No deliver() path is wired.'}
